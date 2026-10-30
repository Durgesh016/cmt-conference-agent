"""Word (.docx) copy of the results: a landscape table with clickable links.

Needs python-docx (pip install python-docx).
"""

from __future__ import annotations

import copy
import re
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .models import Conference

# (header, Conference field, width in cm); widths fill A4 landscape inside 1.27 cm margins.
# The short name isn't a column of its own: it is the end of every CMT link.
COLUMNS = (
    ("Name", "name", 7.0),
    ("Start", "start_date", 1.9),
    ("End", "end_date", 1.9),
    ("Location", "location", 3.5),
    ("CMT link", "cmt_url", 5.5),
    ("Website", "external_url", 3.86),
    ("Keywords", "keywords", 3.5),
)
LINK_FIELDS = {"cmt_url", "external_url"}
HEADER_FILL = "D9E2F3"
LINK_COLOR = "0563C1"
TAB_TITLES = {"my": "My Conferences", "both": "My + All Conferences"}  # anything else: All Conferences
# Characters XML 1.0 forbids (control characters other than tab/newline/CR, and U+FFFE/U+FFFF);
# one of them in a conference name would break the whole file. Mapped to None, translate() drops them.
_NOT_XML = dict.fromkeys([*range(0x00, 0x09), 0x0B, 0x0C, *range(0x0E, 0x20), 0xFFFE, 0xFFFF])


def _xml_safe(text: str) -> str:
    return text.translate(_NOT_XML)


def _link_text(url: str) -> str:
    """What the reader sees: 'https://www.x.org/' -> 'www.x.org' (the link still opens the full URL)."""
    return re.sub(r"^https?://", "", url, flags=re.I).rstrip("/") or url


def _link_target(text: str) -> str | None:
    """A web address Word accepts as a link, or None to keep the text plain.

    Chairs type websites by hand; one with spaces or stray characters as a raw link target can make
    Word call the whole file unreadable, so anything outside the URL character set is %-encoded.
    """
    if not text.lower().startswith(("http://", "https://")) or any(ch.isspace() for ch in text):
        return None
    return quote(text, safe=":/?#[]@!$&'()*+,;=%~")


def _keep_row_on_one_page(row, *, header: bool = False) -> None:
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:cantSplit"))  # schema order: cantSplit before tblHeader
    if header:
        props.append(OxmlElement("w:tblHeader", {qn("w:val"): "true"}))  # repeat the header row on every page


class _Hyperlinks:
    """Adds clickable links, registering each URL once.

    python-docx's own helper rescans every existing link per call, which is far too slow for
    ~20,000 links, so relationship ids are assigned here and each link is a copy of one template.
    """

    def __init__(self, part) -> None:
        self._part = part
        self._ids: dict[str, str] = {}
        props = OxmlElement("w:rPr")  # blue and underlined
        props.append(OxmlElement("w:color", {qn("w:val"): LINK_COLOR}))
        props.append(OxmlElement("w:u", {qn("w:val"): "single"}))
        run = OxmlElement("w:r")
        run.append(props)
        run.append(OxmlElement("w:t", {qn("xml:space"): "preserve"}))
        self._template = OxmlElement("w:hyperlink")
        self._template.append(run)

    def add(self, paragraph, url: str, label: str) -> None:
        """Append to `paragraph` (a w:p element) a link to `url` that shows `label`."""
        r_id = self._ids.get(url)
        if r_id is None:
            r_id = f"rIdLink{len(self._ids) + 1}"
            self._part.rels.add_relationship(RT.HYPERLINK, url, r_id, is_external=True)
            self._ids[url] = r_id
        link = copy.deepcopy(self._template)
        link.set(qn("r:id"), r_id)
        link[0][1].text = label  # the w:t inside the run
        paragraph.append(link)


def _summary(meta: dict) -> str:
    parts = [TAB_TITLES.get(meta.get("tab"), "All Conferences"), f"{meta.get('count', 0):,} conferences"]
    if meta.get("search"):
        parts.append(f"filter: “{meta['search']}”")
    if meta.get("upcoming"):
        parts.append("ending today or later")
    first, last = meta.get("from"), meta.get("to")
    if first and last:
        parts.append(f"starting {first} to {last}")
    elif first:
        parts.append(f"starting from {first}")
    elif last:
        parts.append(f"starting up to {last}")
    parts.append(f"extracted {str(meta.get('extracted_at', ''))[:16].replace('T', ' ')} from {meta.get('source', '')}")
    return "  ·  ".join(parts)


def write_docx(conferences: list[Conference], path: Path, meta: dict) -> Path:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(1.27)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(8.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    document.add_heading("CMT Conferences", level=0)
    summary = document.add_paragraph().add_run(_xml_safe(_summary(meta)))
    summary.font.size = Pt(10)
    document.add_paragraph()

    table = document.add_table(rows=0, cols=len(COLUMNS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, (_, _, width) in zip(table.columns, COLUMNS):
        column.width = Cm(width)  # rows added from here on copy these widths into every cell

    header = table.add_row()
    _keep_row_on_one_page(header, header=True)
    for cell, (title, _, _) in zip(header.cells, COLUMNS):
        shading = OxmlElement("w:shd", {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): HEADER_FILL})
        cell._tc.get_or_add_tcPr().append(shading)
        cell.paragraphs[0].add_run(title).bold = True

    # Every data row has the same shape, so python-docx builds one and the rest are copies of it:
    # its add_row() rescans the whole table each time, which adds up to most of the work at 10,000 rows.
    template = table.add_row()
    _keep_row_on_one_page(template)
    table._tbl.remove(template._tr)
    links = _Hyperlinks(document.part)
    for conference in conferences:
        row = copy.deepcopy(template._tr)
        table._tbl.append(row)
        for cell, (_, field, _) in zip(row.tc_lst, COLUMNS):
            paragraph = cell.p_lst[0]
            value = _xml_safe(getattr(conference, field))
            target = _link_target(value) if field in LINK_FIELDS else None
            if target:
                links.add(paragraph, target, _link_text(value))
            elif value:
                paragraph.add_r().add_t(value)  # values are single-line: whitespace is collapsed upstream

    document.save(str(path))
    return path
